"""
tga_fetcher.py - Fetches and parses unit of competency data from training.gov.au
Strategy: API for metadata, .docm download + robust parsing for content.
"""

import requests
import re
import io
import json
import zipfile
from dataclasses import dataclass, field
from typing import Optional

try:
    from docx import Document as DocxDocument
except ImportError:
    DocxDocument = None


@dataclass
class PerformanceCriteria:
    number: str
    text: str

@dataclass
class Element:
    number: str
    title: str
    performance_criteria: list = field(default_factory=list)

@dataclass
class FoundationSkill:
    skill: str
    description: str
    embedded_pc: str = ""

@dataclass
class KnowledgeItem:
    text: str
    sub_items: list = field(default_factory=list)
    level: int = 0

@dataclass
class UnitData:
    code: str = ""
    title: str = ""
    application: str = ""
    prerequisites: list = field(default_factory=list)
    has_prerequisites: bool = False
    licensing: str = ""
    has_licensing: bool = False
    elements: list = field(default_factory=list)
    foundation_skills: list = field(default_factory=list)
    performance_evidence: list = field(default_factory=list)
    performance_evidence_raw: str = ""
    knowledge_evidence: list = field(default_factory=list)
    knowledge_evidence_raw: str = ""
    assessment_conditions: str = ""
    has_volume_frequency: bool = False
    volume_frequency_detail: str = ""
    has_work_placement: bool = False
    work_placement_detail: str = ""
    has_additional_assessor_reqs: bool = False
    has_simulated_environment: bool = False
    has_required_resources: bool = True
    range_of_conditions: str = ""
    fetch_method: str = "manual"
    fetch_errors: list = field(default_factory=list)

    def is_populated(self):
        return bool(self.code and self.title)


# ── .docm handling ────────────────────────────────────────────────

def docm_to_docx(content: bytes) -> io.BytesIO:
    source = zipfile.ZipFile(io.BytesIO(content))
    output = io.BytesIO()
    with zipfile.ZipFile(output, 'w', zipfile.ZIP_DEFLATED) as dest:
        for item in source.infolist():
            data = source.read(item.filename)
            if item.filename == '[Content_Types].xml':
                data = data.replace(
                    b'application/vnd.ms-word.document.macroEnabled.main+xml',
                    b'application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml'
                )
            if 'vbaProject' in item.filename or 'vbaData' in item.filename:
                continue
            dest.writestr(item, data)
    source.close()
    output.seek(0)
    return output


def open_docm(content: bytes):
    if DocxDocument is None:
        raise ImportError("python-docx not installed")
    try:
        return DocxDocument(io.BytesIO(content))
    except (ValueError, Exception):
        return DocxDocument(docm_to_docx(content))


# ── Network helpers ───────────────────────────────────────────────

def build_download_urls(unit_code: str, package_code: str = None) -> list:
    code = unit_code.upper().strip()
    if package_code:
        pkg_candidates = [package_code]
    else:
        candidates = []
        if len(code) >= 3: candidates.append(code[:3])
        if len(code) >= 2: candidates.append(code[:2])
        if len(code) >= 4: candidates.append(code[:4])
        pkg_candidates = candidates
    
    urls = []
    for pkg in pkg_candidates:
        for release in ["1", "2", "3"]:
            urls.append({"type": "assessment_requirements",
                "url": f"https://training.gov.au/TrainingComponentFiles/{pkg}/{code}_AssessmentRequirements_R{release}.docm",
                "release": release})
            urls.append({"type": "unit",
                "url": f"https://training.gov.au/TrainingComponentFiles/{pkg}/{code}_R{release}.docm",
                "release": release})
    return urls


def try_download_file(url: str, timeout: int = 30) -> Optional[bytes]:
    try:
        print(f"  [download] GET {url}")
        resp = requests.get(url, timeout=timeout, allow_redirects=True)
        print(f"  [download] Status={resp.status_code}, Size={len(resp.content)}")
        if resp.status_code == 200 and len(resp.content) > 500:
            return resp.content
        return None
    except Exception as e:
        print(f"  [download] ERROR: {e}")
        return None


def try_api_endpoint(unit_code: str) -> dict:
    code = unit_code.upper().strip()
    try:
        resp = requests.get(f"https://training.gov.au/api/training/{code}",
                          timeout=10, headers={"Accept": "application/json"})
        if resp.status_code == 200:
            return resp.json()
    except Exception:
        pass
    return {}


# ── Robust .docm parsing ─────────────────────────────────────────

def parse_unit_docm(content: bytes) -> dict:
    """Parse unit .docm: title, application, elements+PCs, foundation skills."""
    doc = open_docm(content)
    data = {"title": "", "application": "", "elements": [], "foundation_skills": []}
    
    print(f"  [parse_unit] Paragraphs: {len(doc.paragraphs)}, Tables: {len(doc.tables)}")
    
    # ── Gather ALL text with source tracking ──
    # Many .docm files put content in tables not paragraphs
    all_lines = []
    for p in doc.paragraphs:
        t = p.text.strip()
        if t:
            all_lines.append({"text": t, "source": "para", "style": p.style.name if p.style else ""})
    
    for ti, table in enumerate(doc.tables):
        for ri, row in enumerate(table.rows):
            for ci, cell in enumerate(row.cells):
                t = cell.text.strip()
                if t:
                    all_lines.append({"text": t, "source": f"table{ti}", "row": ri, "col": ci})
    
    # ── Title ──
    for line in all_lines[:15]:
        t = line["text"]
        if len(t) > 20 and not any(kw in t.lower() for kw in [
            "modification history", "unit of competency", "elements and",
            "application", "prerequisite"
        ]):
            data["title"] = t
            break
    
    # ── Application ──
    # Scan for "Application" label then grab subsequent content
    found_app = False
    app_lines = []
    for i, line in enumerate(all_lines):
        t = line["text"]
        t_lower = t.lower()
        
        if t_lower.startswith("application") and len(t) < 40:
            found_app = True
            continue
        
        if found_app:
            if any(t_lower.startswith(kw) for kw in [
                "element", "prerequisite", "competency field", "unit sector",
                "co-requisite", "foundation skill", "performance criteria",
                "modification history"
            ]):
                break
            # Stop if we hit a very short line that looks like a heading
            if len(t) < 20 and t.endswith(":"):
                break
            app_lines.append(t)
    
    data["application"] = "\n".join(app_lines).strip()
    print(f"  [parse_unit] Application: {len(data['application'])} chars")
    
    # ── Elements and PCs from tables ──
    elements = []
    current_element = None
    
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            if len(cells) < 2:
                continue
            
            left = cells[0]
            right = cells[-1] if len(cells) > 1 else ""
            
            # Skip header rows
            if "element" in left.lower() and "performance" in right.lower():
                continue
            
            el_match = re.match(r'^(\d+)\.\s*(.+)', left)
            
            if el_match and not re.match(r'^\d+\.\d+', left):
                current_element = Element(number=el_match.group(1), title=el_match.group(2).strip())
                elements.append(current_element)
            
            # Check all cells for PCs
            for cell_text in cells:
                for pc_line in cell_text.split("\n"):
                    pc_line = pc_line.strip()
                    pc_match = re.match(r'^(\d+\.\d+)\s+(.+)', pc_line)
                    if pc_match and current_element:
                        # Avoid duplicates
                        pc_num = pc_match.group(1)
                        if not any(pc.number == pc_num for pc in current_element.performance_criteria):
                            current_element.performance_criteria.append(
                                PerformanceCriteria(number=pc_num, text=pc_match.group(2))
                            )
    
    # Fallback: try paragraphs
    if not elements:
        in_elements = False
        for line in all_lines:
            t = line["text"]
            if "elements" in t.lower() and "performance" in t.lower():
                in_elements = True
                continue
            if in_elements:
                if any(t.lower().startswith(kw) for kw in ["foundation", "performance evidence", "knowledge"]):
                    break
                el_match = re.match(r'^(\d+)\.\s+(.+)', t)
                pc_match = re.match(r'^(\d+\.\d+)\s+(.+)', t)
                if el_match and not pc_match:
                    current_element = Element(number=el_match.group(1), title=el_match.group(2))
                    elements.append(current_element)
                elif current_element and pc_match:
                    current_element.performance_criteria.append(
                        PerformanceCriteria(number=pc_match.group(1), text=pc_match.group(2))
                    )
    
    data["elements"] = elements
    total_pcs = sum(len(e.performance_criteria) for e in elements)
    print(f"  [parse_unit] Elements: {len(elements)}, Total PCs: {total_pcs}")
    
    # ── Foundation Skills ──
    foundation_skills = []
    for table in doc.tables:
        header_text = " ".join(c.text.strip().lower() for c in table.rows[0].cells) if table.rows else ""
        
        if "skill" not in header_text and "foundation" not in header_text:
            # Check second row too
            if len(table.rows) > 1:
                row2_text = " ".join(c.text.strip().lower() for c in table.rows[1].cells)
                if "skill" not in row2_text:
                    continue
        
        for row in table.rows[1:]:
            cells = [c.text.strip() for c in row.cells]
            if len(cells) >= 2:
                skill = cells[0]
                desc = cells[1]
                if not skill or not desc or len(skill) > 100:
                    continue
                if "this section" in skill.lower() or skill.lower() in ("skill", ""):
                    continue
                embedded = cells[2] if len(cells) > 2 else ""
                foundation_skills.append(FoundationSkill(skill=skill, description=desc, embedded_pc=embedded))
    
    data["foundation_skills"] = foundation_skills
    print(f"  [parse_unit] Foundation skills: {len(foundation_skills)}")
    
    return data


def parse_assessment_requirements_docm(content: bytes) -> dict:
    """Parse assessment requirements .docm: PE, KE, assessment conditions."""
    doc = open_docm(content)
    data = {
        "performance_evidence": [], "performance_evidence_raw": "",
        "knowledge_evidence": [], "knowledge_evidence_raw": "",
        "assessment_conditions": "",
        "has_simulated": False,
    }
    
    print(f"  [parse_ar] Paragraphs: {len(doc.paragraphs)}, Tables: {len(doc.tables)}")
    
    # Gather all text from paragraphs AND table cells
    all_lines = []
    for p in doc.paragraphs:
        t = p.text.strip()
        if t:
            all_lines.append(t)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                t = cell.text.strip()
                if t and t not in [l for l in all_lines[-5:]]:  # Rough dedup
                    all_lines.append(t)
    
    # Section detection
    current_section = None
    pe_lines, ke_lines, ac_lines = [], [], []
    
    for text in all_lines:
        tl = text.lower().strip()
        
        # Detect section boundaries
        if "performance evidence" in tl and len(text) < 80:
            current_section = "pe"
            continue
        elif "knowledge evidence" in tl and len(text) < 80:
            current_section = "ke"
            continue
        elif "assessment conditions" in tl and len(text) < 80:
            current_section = "ac"
            continue
        elif tl.startswith("foundation skill") and len(text) < 80:
            current_section = None
            continue
        
        if current_section == "pe":
            pe_lines.append(text)
        elif current_section == "ke":
            ke_lines.append(text)
        elif current_section == "ac":
            ac_lines.append(text)
    
    # Clean: remove preamble "The candidate must..." lines
    def clean(lines):
        return [l for l in lines if l and len(l) > 3 
                and not l.lower().startswith("the candidate must")]
    
    data["performance_evidence"] = clean(pe_lines)
    data["performance_evidence_raw"] = "\n".join(pe_lines)
    data["knowledge_evidence"] = [KnowledgeItem(text=l) for l in clean(ke_lines)]
    data["knowledge_evidence_raw"] = "\n".join(ke_lines)
    data["assessment_conditions"] = "\n".join(ac_lines)
    
    # Auto-detect flags
    pe_lower = data["performance_evidence_raw"].lower()
    data["has_volume_frequency"] = any(kw in pe_lower for kw in ["at least", "minimum"])
    if data["has_volume_frequency"]:
        data["volume_frequency_detail"] = data["performance_evidence_raw"]
    
    ac_lower = data["assessment_conditions"].lower()
    data["has_simulated"] = "simulated" in ac_lower
    
    print(f"  [parse_ar] PE={len(data['performance_evidence'])}, KE={len(data['knowledge_evidence'])}, AC={len(data['assessment_conditions'])} chars")
    
    return data


# ── Main fetch ────────────────────────────────────────────────────

def fetch_unit_data(unit_code: str, progress_callback=None) -> UnitData:
    unit = UnitData(code=unit_code.upper().strip())
    
    # Step 1: API
    api_data = None
    if progress_callback:
        progress_callback("Fetching metadata from API...")
    try:
        api_data = try_api_endpoint(unit.code)
        if api_data.get("title"):
            unit.title = api_data["title"]
            unit.fetch_method = "api+docm"
    except Exception as e:
        unit.fetch_errors.append(f"API error: {e}")
    
    # Step 2: Download .docm files
    api_pkg = None
    if api_data and isinstance(api_data, dict):
        parent = api_data.get("parent", {})
        if parent and parent.get("code"):
            api_pkg = parent["code"]
    
    urls = build_download_urls(unit.code, package_code=api_pkg)
    
    ar_content = None
    unit_content = None
    
    if progress_callback:
        progress_callback("Downloading files...")
    
    for u in urls:
        if u["type"] == "unit" and not unit_content:
            c = try_download_file(u["url"])
            if c: unit_content = c
    
    for u in urls:
        if u["type"] == "assessment_requirements" and not ar_content:
            c = try_download_file(u["url"])
            if c: ar_content = c
    
    # Step 3: Parse unit
    if unit_content:
        if progress_callback:
            progress_callback("Parsing unit file...")
        try:
            parsed = parse_unit_docm(unit_content)
            if parsed.get("title") and not unit.title:
                unit.title = parsed["title"]
            if parsed.get("application"):
                unit.application = parsed["application"]
            if parsed.get("elements"):
                unit.elements = parsed["elements"]
            if parsed.get("foundation_skills"):
                unit.foundation_skills = parsed["foundation_skills"]
        except Exception as e:
            unit.fetch_errors.append(f"Unit parse error: {e}")
            import traceback; traceback.print_exc()
    else:
        unit.fetch_errors.append("Could not download unit file.")
    
    # Step 4: Parse assessment requirements
    if ar_content:
        if progress_callback:
            progress_callback("Parsing assessment requirements...")
        try:
            ar = parse_assessment_requirements_docm(ar_content)
            unit.performance_evidence = ar.get("performance_evidence", [])
            unit.performance_evidence_raw = ar.get("performance_evidence_raw", "")
            unit.knowledge_evidence = ar.get("knowledge_evidence", [])
            unit.knowledge_evidence_raw = ar.get("knowledge_evidence_raw", "")
            unit.assessment_conditions = ar.get("assessment_conditions", "")
            unit.has_volume_frequency = ar.get("has_volume_frequency", False)
            unit.volume_frequency_detail = ar.get("volume_frequency_detail", "")
            if ar.get("has_simulated"):
                unit.has_simulated_environment = True
        except Exception as e:
            unit.fetch_errors.append(f"AR parse error: {e}")
            import traceback; traceback.print_exc()
    else:
        unit.fetch_errors.append("Could not download assessment requirements file.")
    
    if not unit.title:
        unit.fetch_method = "manual"
        unit.fetch_errors.append("Could not fetch unit data. Please enter manually.")
    
    return unit
